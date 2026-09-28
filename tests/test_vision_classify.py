"""The vision model, which sees the figure, decides what kind of figure it is.

Classifying from surrounding text sent 12 of 19 figures on the review document
to the flowchart prompt - "no" matched inside "economic", "flow" inside
"Capital Flows" - and a word-bounded version sent every real chart to the
generic prompt instead. The model identified all of them correctly while
complaining it had been asked for a flowchart. So it now names the type on a
first line, and describes the figure by that type's rules.
"""

from __future__ import annotations

import pytest
from docx import Document as DocxDocument

import app.core.vision as vision
from app.core.preprocessing import ADAPTIVE_VISION_PROMPT, _split_type_line, preprocess_docx
from tests.docx_fixtures import picture


def _docx_with_one_figure(tmp_path, caption: str) -> str:
    doc = DocxDocument()
    doc.add_heading("Project economics", level=1)
    doc.add_paragraph(caption)
    doc.add_paragraph().add_run().add_picture(picture(7))
    path = tmp_path / "fig.docx"
    doc.save(path)
    return str(path)


@pytest.fixture
def vision_says(monkeypatch):
    sent = []

    def install(reply: str):
        def describe(items):
            sent.extend(prompt for _, prompt in items)
            return [(reply, None) for _ in items]
        monkeypatch.setattr(vision, "describe_images", describe)
        return sent

    return install


def _figure(result):
    return next(d for d in result.text_docs if d.metadata.get("image_path"))


def test_the_type_the_model_names_becomes_the_content_type(tmp_path, vision_says):
    vision_says("Type: chart\nA bubble scatter plot of LCOE against capacity factor.")
    result = preprocess_docx(_docx_with_one_figure(tmp_path, "Project economics on a purely economic basis."))

    figure = _figure(result)
    assert figure.metadata["content_type"] == "chart"
    assert figure.page_content == "[Chart 1]\nA bubble scatter plot of LCOE against capacity factor."
    assert [v.content_type for v in result.visual_elements if v.image_path] == ["chart"]


def test_every_docx_figure_gets_the_adaptive_prompt(tmp_path, vision_says):
    sent = vision_says("Type: image\nA logo.")
    preprocess_docx(_docx_with_one_figure(tmp_path, "Capital Flows to 2026"))

    assert sent == [ADAPTIVE_VISION_PROMPT]


def test_without_a_type_line_the_figure_is_an_image_and_nothing_is_cut(tmp_path, vision_says):
    vision_says("A wind farm at dusk.")
    figure = _figure(preprocess_docx(_docx_with_one_figure(tmp_path, "Site photography")))

    assert figure.metadata["content_type"] == "image"
    assert figure.page_content == "[Image 1]\nA wind farm at dusk."


@pytest.mark.parametrize("line, expected", [
    ("Type: chart", "chart"),
    ("**Type:** Flowchart", "flowchart"),
    ("type: DIAGRAM", "diagram"),
    ("Type: table", "table"),
    ("Type: image", "image"),
])
def test_the_type_line_is_read_in_the_forms_models_write_it(line, expected):
    kind, rest = _split_type_line(f"{line}\nThe description.")
    assert (kind, rest) == (expected, "The description.")


def test_an_unknown_type_is_not_trusted():
    kind, rest = _split_type_line("Type: spaceship\nThe description.")
    assert (kind, rest) == (None, "Type: spaceship\nThe description.")


def test_the_prompt_keeps_every_type_specific_instruction():
    """One prompt, not five - but none of the five sets of rules is lost."""
    for rule in ("axes", "decision points", "relationships", "Markdown", "transcribed exactly"):
        assert rule in ADAPTIVE_VISION_PROMPT
