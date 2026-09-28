"""Tables and figures carry the heading they sit under.

The DOCX was read in three separate passes - paragraphs, `docx.tables`, image
relationships - and only the paragraph text went through the header splitter.
No table or figure ever learned its heading, so a question worded like the
heading ("attractiveness ranking") matched neither retriever. Measured on the
review document: 0 of 34 tables had a section breadcrumb.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as DocxDocument

import app.core.vision as vision
from app.core.preprocessing import preprocess_docx
from tests.docx_fixtures import picture


@pytest.fixture
def described(monkeypatch):
    """Vision is an Azure call; describe each image by its filename instead."""
    monkeypatch.setattr(
        vision, "describe_images",
        lambda items: [(f"described {Path(path).name}", None) for path, _ in items],
    )


def _build(tmp_path) -> str:
    doc = DocxDocument()
    doc.add_heading("Overview", level=1)
    doc.add_paragraph("Intro text for the overview.")
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "first table"

    doc.add_heading("Composite ranking", level=2)
    doc.add_table(rows=2, cols=2).cell(0, 0).text = "second table"
    doc.add_paragraph().add_run().add_picture(picture(1))          # A

    doc.add_heading("Risks", level=1)
    doc.add_heading("Supply", level=2)
    third = doc.add_table(rows=2, cols=2)
    third.cell(0, 0).text = "third table"
    third.cell(1, 1).paragraphs[0].add_run().add_picture(picture(2))  # B, in a cell
    doc.add_paragraph().add_run().add_picture(picture(1))          # A again

    doc.part.get_or_add_image(picture(3))                           # C, never placed

    path = tmp_path / "walk.docx"
    doc.save(path)
    return str(path)


def _by_type(result, content_type):
    return [d for d in result.text_docs if d.metadata.get("content_type") == content_type]


def test_each_table_carries_the_heading_path_above_it(tmp_path, described):
    tables = _by_type(preprocess_docx(_build(tmp_path)), "table")

    assert [t.metadata.get("section_header") for t in tables] == [
        "Overview", "Overview > Composite ranking", "Risks > Supply",
    ]
    assert [t.page_content.splitlines()[0] for t in tables] == [
        "[Table 1]", "[Table 2]", "[Table 3]",
    ]


def test_a_figure_carries_the_heading_it_sits_under(tmp_path, described):
    result = preprocess_docx(_build(tmp_path))
    figures = [d for d in result.text_docs if d.metadata.get("image_path")]
    sections = {Path(d.metadata["image_path"]).name: d.metadata.get("section_header")
                for d in figures}

    assert sections == {
        "walk_img1.png": "Overview > Composite ranking",
        "walk_img2.png": "Risks > Supply",
        "walk_img3.png": None,
    }


def test_a_picture_inside_a_table_cell_is_emitted_once(tmp_path, described):
    """Also covers the same picture placed twice: one figure, not two."""
    result = preprocess_docx(_build(tmp_path))
    figures = [d for d in result.text_docs if d.metadata.get("image_path")]

    assert len(figures) == 3


def test_an_image_the_body_never_places_is_still_extracted(tmp_path, described):
    result = preprocess_docx(_build(tmp_path))
    names = [Path(d.metadata["image_path"]).name
             for d in result.text_docs if d.metadata.get("image_path")]

    assert "walk_img3.png" in names


def test_the_text_stream_comes_first_with_its_headings(tmp_path, described):
    text = preprocess_docx(_build(tmp_path)).text_docs[0]

    assert text.metadata["content_type"] == "text"
    assert "# Overview" in text.page_content
    assert "## Composite ranking" in text.page_content
    assert "Intro text for the overview." in text.page_content
