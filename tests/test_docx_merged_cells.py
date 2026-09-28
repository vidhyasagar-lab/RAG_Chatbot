"""A merged cell is written once, not once per column it spans.

python-docx yields the same cell for every grid column a horizontal merge
covers, so Table 12 on the review document rendered its title four times.
Blanking the repeats keeps the column count, so the Markdown stays aligned.
Identity, not text, decides: two separate cells that both read "I" are real
data (RACI tables do this) and must stay.
"""

from __future__ import annotations

from docx import Document as DocxDocument

from app.core.preprocessing import _docx_table_to_markdown


def _table():
    table = DocxDocument().add_table(rows=5, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = "Title"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "Low"
    table.cell(1, 2).text = "High"
    table.cell(2, 0).merge(table.cell(3, 0)).text = "V"
    for col, text in enumerate(["I", "I", "x"]):
        table.cell(4, col).text = text
    return table


def _lines():
    return _docx_table_to_markdown(_table()).splitlines()


def test_a_horizontal_merge_is_written_once():
    assert _lines()[0] == "| Title |  |  |"


def test_a_partial_merge_keeps_its_neighbour():
    assert _lines()[2] == "| Low |  | High |"   # [1] is the --- separator


def test_the_column_count_is_unchanged():
    assert all(line.count("|") == 4 for line in _lines())


def test_a_vertical_merge_still_repeats_down_the_column():
    """Each retrieved row should read on its own."""
    lines = _lines()
    assert lines[3].startswith("| V |")
    assert lines[4].startswith("| V |")


def test_equal_text_in_separate_cells_is_kept():
    assert _lines()[5] == "| I | I | x |"
